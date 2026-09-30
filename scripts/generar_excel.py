"""Genera datos/Registro_Terra_PSCJ.xlsx desde data/propiedades.json y el historial de git."""
import json, os, subprocess
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = json.load(open(os.path.join(ROOT, "data", "propiedades.json"), encoding="utf-8"))
try:
    out = subprocess.run(["git", "log", "-n", "300", "--date=format:%Y-%m-%d %H:%M", "--pretty=format:%ad|%an|%s"], cwd=ROOT, capture_output=True, text=True, check=True).stdout
    H = [l.split("|", 2) for l in out.splitlines() if l.count("|") >= 2]
except Exception:
    H = []

wb = Workbook(); ws = wb.active; ws.title = "Resumen"
hf, hb = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0D2A4D")
def sheet(name, cols, rows, widths, money=None, dv=None):
    s = wb.create_sheet(name); s.append(cols)
    for c in s[1]: c.font, c.fill, c.alignment = hf, hb, Alignment(vertical="center", wrap_text=True)
    for r in rows: s.append(list(r))
    for i, w in enumerate(widths, 1): s.column_dimensions[get_column_letter(i)].width = w
    s.freeze_panes = "A2"; s.auto_filter.ref = s.dimensions; s.row_dimensions[1].height = 24
    if money:
        for row in s.iter_rows(min_row=2, min_col=money, max_col=money):
            for c in row: c.number_format = '"$"#,##0'
    if dv:
        v = DataValidation(type="list", formula1='"' + ",".join(dv[1]) + '"', allow_blank=True); s.add_data_validation(v); v.add(f"{dv[0]}2:{dv[0]}2000")
sheet("Propiedades", ["ID", "Título", "Tipo", "Ubicación", "Área m²", "Precio COP", "Habitaciones", "Descripción", "Estado", "Publicada"],
      [(p["id"], p["titulo"], p["tipo"], p.get("ubicacion"), p.get("area_m2"), p.get("precio"), p.get("habitaciones"), p.get("descripcion"), p.get("estado"), "Sí" if p.get("publicada") else "No") for p in P],
      [6, 38, 12, 24, 10, 16, 13, 46, 13, 11], money=6, dv=("I", ["Disponible", "Reservada", "Vendida"]))
sheet("Historial", ["Fecha", "Quién", "Cambio"], H, [18, 24, 70])
ws["A1"], ws["A1"].font = "Terra PSCJ · Registro de gestión", Font(bold=True, size=16, color="0D2A4D")
ws["A2"] = "Generado: " + datetime.now().strftime("%Y-%m-%d %H:%M")
rows = [("Propiedades registradas", "=COUNTA(Propiedades!A2:A2000)"), ("Disponibles", '=COUNTIF(Propiedades!I2:I2000,"Disponible")'), ("Reservadas", '=COUNTIF(Propiedades!I2:I2000,"Reservada")'),
        ("Vendidas", '=COUNTIF(Propiedades!I2:I2000,"Vendida")'), ("Valor del inventario disponible (COP)", '=SUMIF(Propiedades!I2:I2000,"Disponible",Propiedades!F2:F2000)'), ("Cambios registrados", "=COUNTA(Historial!A2:A2000)")]
for i, (a, b) in enumerate(rows, 4): ws.cell(i, 1, a).font = Font(bold=True); ws.cell(i, 2, b)
ws["B8"].number_format = '"$"#,##0'; ws.column_dimensions["A"].width = 42; ws.column_dimensions["B"].width = 20
os.makedirs(os.path.join(ROOT, "datos"), exist_ok=True)
wb.save(os.path.join(ROOT, "datos", "Registro_Terra_PSCJ.xlsx")); print("Excel generado")
